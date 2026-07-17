#!/usr/bin/env python3
"""
Research URL Fetcher
Reads URLS_JSON and SESSION_ID from environment variables.

Outputs:
  research-logs/fetch-log.xlsx        -- cumulative log, one row per attempt
  fetched-content/<session_id>/       -- extracted plain text, one file per URL
"""

import io
import json
import os
import re
import sys
import time
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse

import httpx
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

try:
    import pdfplumber
    HAS_PDF = True
except ImportError:
    HAS_PDF = False
    print("Warning: pdfplumber not available; PDF extraction disabled.")

# ── Palette (workspace palette) ───────────────────────────────────────────────
NAVY  = "002060"
BLUE  = "0070C0"
TEAL  = "006B6B"
WHITE = "FFFFFF"
INK   = "1A1C22"

STATUS_STYLE = {
    "Success": ("C6EFCE", "1A5E38"),  # green
    "Partial": ("FFEB9C", "7B4F00"),  # amber
    "Failed":  ("FFC7CE", "9C0006"),  # red
    "Timeout": ("FFD9B0", "7B3200"),  # orange
    "Error":   ("FFC7CE", "9C0006"),  # red
}

# (header label, column width)
COLUMNS = [
    ("Timestamp (UTC)",  22),
    ("Session",          18),
    ("Document Name",    40),
    ("URL",              55),
    ("Status",           12),
    ("HTTP Code",        10),
    ("Content Type",     14),
    ("Words Extracted",  16),
    ("Summary",          65),
    ("Duration (s)",     12),
]

LOG_PATH    = Path("research-logs/fetch-log.xlsx")
CONTENT_DIR = Path("fetched-content")


# ── Text extraction ───────────────────────────────────────────────────────────

def _title_from_url(url: str) -> str:
    path = urlparse(url).path.rstrip("/").split("/")[-1]
    return re.sub(r"[_-]", " ", path).strip()[:100] or url[:80]


def extract_html(raw: bytes, url: str) -> tuple:
    """Returns (title, full_text, word_count)."""
    try:
        soup = BeautifulSoup(raw, "lxml")
        for tag in soup(["script", "style", "nav", "header", "footer", "aside", "noscript"]):
            tag.decompose()
        title_tag = soup.find("title")
        title = title_tag.get_text(strip=True)[:120] if title_tag else ""
        if not title:
            h1 = soup.find("h1")
            title = h1.get_text(strip=True)[:120] if h1 else _title_from_url(url)
        text = soup.get_text(separator=" ", strip=True)
        text = re.sub(r" {3,}", "  ", text)
        return title, text, len(text.split())
    except Exception as exc:
        return _title_from_url(url), f"[HTML error: {exc}]", 0


def extract_pdf(raw: bytes, url: str) -> tuple:
    """Returns (title, full_text, word_count)."""
    if not HAS_PDF:
        return _title_from_url(url), "[pdfplumber not installed]", 0
    try:
        with pdfplumber.open(io.BytesIO(raw)) as pdf:
            meta   = pdf.metadata or {}
            title  = str(meta.get("Title", "")).strip()[:120]
            pages  = [p.extract_text() or "" for p in pdf.pages[:60]]
        if not title and pages:
            title = pages[0].strip().splitlines()[0][:120]
        title = title or _title_from_url(url)
        text  = "\n\n".join(pages)
        return title, text, len(text.split())
    except Exception as exc:
        return _title_from_url(url), f"[PDF error: {exc}]", 0


# ── HTTP fetch ────────────────────────────────────────────────────────────────

def fetch_url(url: str, timeout: int = 30) -> dict:
    result = {
        "url": url, "title": "", "status": "Failed",
        "http_code": "", "content_type": "", "words": 0,
        "summary": "", "full_text": "", "duration": 0.0,
    }
    t0 = time.time()
    headers = {
        "User-Agent": (
            "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
            "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
        ),
        "Accept": "text/html,application/pdf,application/xhtml+xml,*/*;q=0.8",
        "Accept-Language": "en-US,en;q=0.9",
    }
    try:
        with httpx.Client(follow_redirects=True, timeout=timeout, verify=True) as client:
            r = client.get(url, headers=headers)

        result["http_code"] = r.status_code
        result["duration"]  = round(time.time() - t0, 2)

        if r.status_code == 200:
            ct     = r.headers.get("content-type", "").lower()
            is_pdf = "pdf" in ct
            result["content_type"] = "PDF" if is_pdf else "HTML"

            if is_pdf:
                title, full_text, words = extract_pdf(r.content, url)
            else:
                title, full_text, words = extract_html(r.content, url)

            result["title"]     = title
            result["full_text"] = full_text
            result["words"]     = words
            result["summary"]   = " ".join(full_text.split()[:55])[:300]
            result["status"]    = "Success" if words > 100 else "Partial"
        else:
            reason = getattr(r, "reason_phrase", "") or ""
            result["summary"] = f"HTTP {r.status_code} {reason}".strip()

    except httpx.TimeoutException:
        result["status"]   = "Timeout"
        result["duration"] = round(time.time() - t0, 2)
        result["summary"]  = f"No response within {timeout}s"
    except Exception as exc:
        result["status"]   = "Error"
        result["duration"] = round(time.time() - t0, 2)
        result["summary"]  = str(exc)[:200]

    if not result["title"]:
        result["title"] = _title_from_url(url)

    return result


# ── Excel helpers ─────────────────────────────────────────────────────────────

def _border(color="DDDDDD"):
    s = Side(style="thin", color=color)
    return Border(left=s, right=s, top=s, bottom=s)


def _build_header(ws):
    hdr_fill   = PatternFill("solid", fgColor=NAVY)
    hdr_font   = Font(name="Arial", bold=True, size=10, color=WHITE)
    hdr_align  = Alignment(horizontal="center", vertical="center", wrap_text=True)
    hdr_border = _border("888888")

    for col, (label, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.alignment = hdr_align
        cell.border    = hdr_border
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def load_or_create_workbook():
    if LOG_PATH.exists():
        try:
            wb = openpyxl.load_workbook(LOG_PATH)
            return wb, wb.active
        except Exception as exc:
            print(f"Warning: existing log unreadable ({exc}); starting fresh.")
            LOG_PATH.rename(LOG_PATH.with_suffix(".bak.xlsx"))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fetch Log"
    _build_header(ws)
    return wb, ws


def append_row(ws, result: dict, session_id: str, ts: str):
    fill_hex, font_hex = STATUS_STYLE.get(result["status"], ("F7F8FA", INK))
    row_fill = PatternFill("solid", fgColor=fill_hex)
    row_n    = ws.max_row + 1

    CENTER = Alignment(horizontal="center", vertical="top")
    WRAP   = Alignment(horizontal="left",   vertical="top", wrap_text=True)

    values = [
        ts,
        session_id,
        (result.get("title") or "")[:120],
        result["url"],           # becomes hyperlink
        result["status"],
        str(result.get("http_code", "")),
        result.get("content_type", ""),
        result.get("words", 0),
        result.get("summary", ""),
        result.get("duration", ""),
    ]
    wrap_cols  = {3, 4, 9}   # Document Name, URL, Summary
    center_cols = {1, 2, 5, 6, 7, 8, 10}

    for col, value in enumerate(values, start=1):
        cell           = ws.cell(row=row_n, column=col, value=value)
        cell.fill      = row_fill
        cell.border    = _border()
        cell.alignment = WRAP if col in wrap_cols else CENTER

        if col == 4:  # URL column — hyperlink
            cell.font      = Font(name="Arial", size=9, color=BLUE, underline="single")
            cell.hyperlink = result["url"]
        else:
            cell.font = Font(name="Arial", size=9, color=font_hex)

    ws.row_dimensions[row_n].height = 55


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    urls_json  = os.environ.get("URLS_JSON", "").strip()
    raw_sid    = os.environ.get("SESSION_ID", "research")
    session_id = re.sub(r"[^a-zA-Z0-9_-]", "-", raw_sid)

    if not urls_json:
        print("Error: URLS_JSON environment variable is required.", file=sys.stderr)
        sys.exit(1)

    try:
        urls = json.loads(urls_json)
        if isinstance(urls, str):        # handle double-encoded input
            urls = json.loads(urls)
    except json.JSONDecodeError as exc:
        print(f"Error: URLS_JSON is not valid JSON: {exc}", file=sys.stderr)
        sys.exit(1)

    if not isinstance(urls, list):
        urls = [str(urls)]

    print(f"Session: '{session_id}'  |  URLs: {len(urls)}\n{'─' * 60}")

    LOG_PATH.parent.mkdir(parents=True, exist_ok=True)
    out_dir = CONTENT_DIR / session_id
    out_dir.mkdir(parents=True, exist_ok=True)

    wb, ws = load_or_create_workbook()

    failed = 0
    for i, url in enumerate(urls, start=1):
        ts     = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
        print(f"[{i:>2}/{len(urls)}] {url}")
        result = fetch_url(url)

        if result["full_text"] and result["status"] in ("Success", "Partial"):
            slug      = re.sub(r"[^a-z0-9]+", "-",
                               (urlparse(url).netloc + urlparse(url).path).lower())
            slug      = slug.strip("-")[:80]
            text_path = out_dir / f"{slug}.txt"
            text_path.write_text(result["full_text"], encoding="utf-8", errors="replace")
            print(f"        {result['status']} | HTTP {result['http_code']} "
                  f"| {result['words']:,} words | {result['duration']}s "
                  f"| -> {text_path.name}")
        else:
            failed += 1
            print(f"        {result['status']} | HTTP {result.get('http_code', '')} "
                  f"| {result['duration']}s | {result['summary'][:80]}")

        append_row(ws, result, session_id, ts)

    wb.save(LOG_PATH)
    total_rows = ws.max_row - 1
    print(f"\n{'─' * 60}")
    print(f"Log saved: {LOG_PATH}  ({total_rows} total entries)")
    print(f"Success/Partial: {len(urls) - failed}/{len(urls)} | Failed: {failed}/{len(urls)}")
    if failed == len(urls):
        sys.exit(1)   # signal CI failure if everything failed


if __name__ == "__main__":
    main()
